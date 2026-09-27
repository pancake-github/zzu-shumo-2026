"""生成可编辑的中文三线表；数字和西文使用 Times New Roman。"""  # 模块用途说明。
from __future__ import annotations  # 允许使用前向类型注解。
from collections.abc import Mapping, Sequence  # 引入可迭代表格和富文本结构类型。
from dataclasses import dataclass  # 用数据类描述独立文字片段。
import re  # 支持普通表序及章节式表序。
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # 引入表格对齐常量。
from docx.enum.text import WD_ALIGN_PARAGRAPH  # 引入段落对齐常量。
from docx.oxml import OxmlElement  # 创建三线表所需的原生 Word XML 元素。
from docx.oxml.ns import qn  # 将带前缀的 XML 名称转成完整名称。
from docx.shared import Cm, Pt, RGBColor  # 引入尺寸和颜色工具。

@dataclass(frozen=True)  # 每个富文本片段独立保存排版信息。
class RunSpec:  # 用于中文名称、物理量、下标和单位的分段排版。
    text: str  # 文字内容直接使用 Unicode，不在此处解释 LaTeX。
    italic: bool = False  # 量符号设为斜体，单位和普通文字保持正体。
    bold: bool = False  # 默认不加粗，可按文档需求显式指定。
    subscript: bool = False  # 描述性下标可用正体片段并开启本项。
    superscript: bool = False  # 指数可用独立片段并开启本项。

CellContent = str | Sequence[RunSpec | Mapping[str, object]]  # 一个单元格可为纯文本或多个格式片段。

def _set_font_mapping(properties, *, size_pt: float, chinese_font: str = "宋体", latin_font: str = "Times New Roman") -> None:  # 写入字体映射，避免 Word 主题字体覆盖。
    fonts = properties.find(qn("w:rFonts"))  # 查找已有字体设置。
    if fonts is None:  # 没有字体设置时创建字体元素。
        fonts = OxmlElement("w:rFonts")  # 创建原生 Word 字体属性。
        properties.insert(0, fonts)  # 将字体属性置于字符属性前部。
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):  # 清除可能优先使用的主题字体。
        fonts.attrib.pop(qn(f"w:{attribute}"), None)  # 删除现有主题属性而不影响其余设置。
    for attribute in ("ascii", "hAnsi", "cs"):  # 西文、数字、希腊字符优先走西文字体。
        fonts.set(qn(f"w:{attribute}"), latin_font)
    fonts.set(qn("w:eastAsia"), chinese_font)  # 实际可用性需渲染核验。
    for element_name in ("sz", "szCs"):  # 同时设置普通字符和复杂文字字号。
        element = properties.find(qn(f"w:{element_name}"))  # 查找原字号属性。
        if element is None:  # 缺少字号属性时补齐。
            element = OxmlElement(f"w:{element_name}")  # 创建字号节点。
            properties.append(element)  # 插入字符属性。
        element.set(qn("w:val"), str(round(size_pt * 2)))  # Word 字号使用半磅整数。

def configure_document_fonts(document, *, body_size_pt: float = 11.0, chinese_font: str = "宋体", latin_font: str = "Times New Roman") -> None:  # 设置新建文档字体；修订既有文档时不调用此全局函数。
    if body_size_pt <= 0:  # 拒绝不可用的字号。
        raise ValueError("正文字号必须大于零。")  # 明确报告参数问题。
    for style_name in ("Normal", "Caption"):  # 表格正文和表题采用一致的中西文映射。
        style = document.styles[style_name]  # 取得目标段落样式。
        style.font.name = latin_font
        style.font.size = Pt(body_size_pt)  # 设置基础字号。
        style.font.color.rgb = RGBColor(0, 0, 0)  # 使用黑色文字。
        _set_font_mapping(style.element.get_or_add_rPr(), size_pt=body_size_pt, chinese_font=chinese_font, latin_font=latin_font)

def add_rich_text(paragraph, content: CellContent, *, size_pt: float = 11.0, chinese_font: str = "宋体", latin_font: str = "Times New Roman") -> None:  # 向已有段落写入可编辑富文本。
    specs = [RunSpec(content)] if isinstance(content, str) else content  # 统一纯文本和富文本输入。
    for item in specs:  # 按原顺序插入片段，保持物理量和单位的边界。
        spec = RunSpec(**item) if isinstance(item, Mapping) else item  # 同时支持字典和数据类写法。
        if not isinstance(spec, RunSpec) or not isinstance(spec.text, str):  # 检查文本片段类型。
            raise TypeError("单元格须为字符串或由 RunSpec/字典组成的序列。")  # 不自动猜测未知输入。
        if spec.subscript and spec.superscript:  # 同一片段不能同时位于上下标。
            raise ValueError("同一个文字片段不能同时设置上标和下标。")  # 提醒使用分开的片段。
        run = paragraph.add_run(spec.text)  # 写入原生可编辑文字。
        run.font.name = latin_font
        run.font.size = Pt(size_pt)  # 设置片段字号。
        run.font.color.rgb = RGBColor(0, 0, 0)  # 使用黑色字符。
        run.italic = spec.italic  # 只对显式标记的量符号使用斜体。
        run.bold = spec.bold  # 按片段要求设置粗体。
        if spec.subscript:  # 单独处理下标，避免上下标开关互相覆盖。
            run.font.subscript = True  # 设置 Word 原生下标。
        elif spec.superscript:  # 没有下标时才处理上标。
            run.font.superscript = True  # 设置 Word 原生上标。
        _set_font_mapping(run._r.get_or_add_rPr(), size_pt=size_pt, chinese_font=chinese_font, latin_font=latin_font)

def _set_border(container, side: str, *, visible: bool, size: int = 8) -> None:  # 在表格或单元格边框容器中写入指定边。
    border = OxmlElement(f"w:{side}")  # 创建指定方向的边框元素。
    border.set(qn("w:val"), "single" if visible else "nil")  # 显示单实线或明确禁用边框。
    if visible:  # 可见边框统一为黑色。
        border.set(qn("w:sz"), str(size))  # 线宽单位为八分之一磅。
        border.set(qn("w:color"), "000000")  # 设置黑色边线。
        border.set(qn("w:space"), "0")  # 不附加线外间距。
    container.append(border)  # 将边线加入对应边框容器。

def _prepare_cell(cell, *, top: bool, middle: bool, bottom: bool) -> None:  # 清除网格并按行位置设置三条横线。
    properties = cell._tc.get_or_add_tcPr()  # 取得单元格属性。
    for name in ("tcBorders", "shd", "tcMar"):  # 防止模板残留边线、底色或紧凑边距。
        for element in list(properties.findall(qn(f"w:{name}"))):  # 收集该类现有属性。
            properties.remove(element)  # 删除旧属性以便统一重建。
    borders = OxmlElement("w:tcBorders")  # 新建单元格边线容器。
    for side in ("top", "left", "bottom", "right", "insideH", "insideV", "tl2br", "tr2bl"):  # 显式关闭竖线、内部线和对角线。
        visible = (side == "top" and top) or (side == "bottom" and (middle or bottom))  # 仅表顶、表头下方和表底显示线条。
        _set_border(borders, side, visible=visible, size=6 if middle and side == "bottom" else 10)  # 表头分隔线细于上下边线。
    properties.append(borders)  # 应用三线规则。
    shading = OxmlElement("w:shd")  # 显式使用白底。
    shading.set(qn("w:val"), "clear")  # 清除图案填充。
    shading.set(qn("w:fill"), "FFFFFF")  # 设置白色背景。
    properties.append(shading)  # 写入底色属性。
    margins = OxmlElement("w:tcMar")  # 设置单元格留白。
    for side, width in (("top", 90), ("bottom", 90), ("left", 110), ("right", 110)):  # 上下和左右使用适度边距。
        margin = OxmlElement(f"w:{side}")  # 创建对应方向的边距。
        margin.set(qn("w:w"), str(width))  # 边距单位为二十分之一磅。
        margin.set(qn("w:type"), "dxa")  # 使用绝对距离，避免缩放误差。
        margins.append(margin)  # 加入边距容器。
    properties.append(margins)  # 写入单元格边距。
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER  # 单元格文字垂直居中。

def add_three_line_table(document, *, number: int | str, title: str, headers: Sequence[CellContent], rows: Sequence[Sequence[CellContent]], widths_cm: Sequence[float] | None = None, alignments: Sequence[str] | None = None, size_pt: float = 10.5, chinese_font: str = "宋体", latin_font: str = "Times New Roman"):
    """新建可编辑三线表；字体和编号服从调用者的模板，不修改已有表格。"""
    if isinstance(number, bool) or not re.fullmatch(r'[1-9]\d*(?:[.-][1-9]\d*)?', str(number)):
        raise ValueError("表序须为正整数或章节式编号，如 2.1。")
    if not title.strip() or "\n" in title or "\r" in title:  # 简明表题保持单段结构。
        raise ValueError("表题不能为空或包含换行。")  # 要求调用者提供实际中文表题。
    if not headers or not rows:  # 空表无法形成独立的三条横线。
        raise ValueError("请提供表头和至少一行真实数据；没有数据时不生成占位表。")  # 阻止空数据的占位输出。
    count = len(headers)  # 记录表格列数。
    if any(len(row) != count for row in rows):  # 每行应与表头保持一致。
        raise ValueError("所有数据行的列数必须与表头相同。")  # 避免静默丢失字段。
    section = document.sections[-1]  # 按当前节的页面宽度计算可用区域。
    available_cm = (section.page_width - section.left_margin - section.right_margin) / 360000  # 将 Word 长度换算为厘米。
    widths = list(widths_cm) if widths_cm is not None else [available_cm / count] * count  # 未指定时等宽，正式表格宜按内容给定宽度。
    if len(widths) != count or any(width <= 0 for width in widths) or sum(widths) > available_cm + 0.01:  # 检查列宽数目、正值与页面边界。
        raise ValueError("列宽须为每列一个正数，合计不得超过当前节的正文宽度。")  # 防止表格越过页边距。
    positions = list(alignments) if alignments is not None else ["center"] * count  # 数值表默认居中，可逐列指定。
    position_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}  # 将可读名称映射至 Word 对齐方式。
    if len(positions) != count or any(position not in position_map for position in positions):  # 验证每列对齐设置。
        raise ValueError("每列对齐方式只能是 left、center 或 right。")  # 拒绝无法解释的参数。
    if size_pt <= 0:  # 字号必须可显示。
        raise ValueError("表格字号必须大于零。")  # 明确说明错误。
    caption = document.add_paragraph(style="Caption")  # 在表格前插入可编辑表题。
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 表题居中。
    caption.paragraph_format.keep_with_next = True  # 表题与紧随其后的表格保持同页。
    caption.paragraph_format.space_before = Pt(8)  # 为表格之前留白。
    caption.paragraph_format.space_after = Pt(5)  # 表题与表格之间保留适度空白。
    add_rich_text(caption, f"表{number}　{title}", size_pt=size_pt, chinese_font=chinese_font, latin_font=latin_font)
    table = document.add_table(rows=len(rows) + 1, cols=count)  # 创建表头和数据行。
    table.alignment = WD_TABLE_ALIGNMENT.CENTER  # 表格整体居中。
    table.autofit = False  # 保持调用者给定的列宽。
    properties = table._tbl.tblPr  # 取得表级属性。
    inherited_style = properties.find(qn("w:tblStyle"))  # 检查继承的表格样式。
    if inherited_style is not None:  # 避免模板样式重新加回网格线。
        properties.remove(inherited_style)  # 移除当前表的样式引用。
    table_borders = OxmlElement("w:tblBorders")  # 表级边框全部关闭，三条横线由单元格明确设置。
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):  # 逐项禁用表级外框和网格。
        _set_border(table_borders, side, visible=False)  # 避免隐式默认边框。
    properties.append(table_borders)  # 写入表级边框规则。
    for column, width in zip(table.columns, widths):  # 同步 Word 的列网格宽度。
        column.width = Cm(width)  # 按厘米设置各列宽度。
    all_rows = [headers, *rows]  # 将表头和数据按显示顺序排列。
    for row_index, (row, values) in enumerate(zip(table.rows, all_rows)):  # 按行插入数据并设置分页属性。
        row_properties = row._tr.get_or_add_trPr()  # 取得行级属性。
        row_properties.append(OxmlElement("w:cantSplit"))  # 禁止同一行跨页拆分。
        if row_index == 0:  # 首行同时作为可重复的表头。
            row_properties.append(OxmlElement("w:tblHeader"))  # 跨页时重复列标题。
        for column_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):  # 按列写入正文和格式。
            cell.width = Cm(width)  # 同步单元格宽度和表网格宽度。
            _prepare_cell(cell, top=row_index == 0, middle=row_index == 0, bottom=row_index == len(all_rows) - 1)  # 仅在表顶、表头下方和表底画线。
            paragraph = cell.paragraphs[0]  # 使用单元格中的初始段落。
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else position_map[positions[column_index]]  # 表头居中，正文逐列对齐。
            paragraph.paragraph_format.space_before = Pt(0)  # 单元格中不叠加段前距。
            paragraph.paragraph_format.space_after = Pt(0)  # 单元格中不叠加段后距。
            paragraph.paragraph_format.line_spacing = 1.15  # 让上下标和希腊字符有足够行高。
            paragraph.paragraph_format.keep_with_next = row_index == 0  # 表头与首行数据保持连续，正文允许跨页。
            add_rich_text(paragraph, value, size_pt=size_pt, chinese_font=chinese_font, latin_font=latin_font)
    return table  # 返回可继续编辑的 python-docx 表格对象。
